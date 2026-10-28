"""Traži preostali srpski tekst u izlazima:  python tools/scan_sr.py DIR"""
import sys, os, re, glob, collections
d = sys.argv[1]
SRW = re.compile(r"[čćžšđČĆŽŠĐ]|\b(je|su|za|od|do|sa|ili|kao|koji|preglednost\w*|brzina|trase?|dionic\w*|vozil\w*|proračun\w*|tabel\w*|slika|nagib\w*|osovin\w*|uslov\w*|dobro|loše|pravac|luk|lukov\w*|stacionaž\w*|ocjena|smjer\w*|traka|zraka|prijelaznica|geometrija|ukupno|sekcija|dužina|zona|kolovoz\w*|krivina|elementi?|ulaz|prilog\w*)\b", re.I)
hits = collections.Counter()
def chk(src, t):
    t = str(t)
    if SRW.search(t): hits[(src, t.strip()[:100])] += 1
from docx import Document
for f in glob.glob(d + "/*.docx"):
    doc = Document(f)
    for p in doc.paragraphs: chk(os.path.basename(f)[:12], p.text)
    for tb in doc.tables:
        for r in tb.rows:
            for c in r.cells: chk("docx-tbl", c.text)
import openpyxl
for f in glob.glob(d + "/*.xlsx"):
    wb = openpyxl.load_workbook(f)
    for ws in wb:
        chk("sheetname", ws.title)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str): chk(os.path.basename(f)[:14] + ":" + ws.title, c.value)
for f in glob.glob(d + "/*.dxf"):
    for l in open(f, encoding="utf-8", errors="ignore").read().split("\n"):
        chk("dxf", l)
for (s, t), n in sorted(hits.items()): print(f"{s} | {t}")
print(len(hits), "hits")
