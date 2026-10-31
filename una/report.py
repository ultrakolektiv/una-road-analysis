"""Generator Word elaborata (Markdown → pandoc → python-docx dorada)."""
from .i18n import _T, _TF, tr_obj, cell as _cell, dec as _dec
import math, os, re, shutil, subprocess, tempfile
import numpy as np
from .texts import SYM, LIT
from .engine import TIPOVI
from . import engine

HERE = os.path.dirname(os.path.abspath(__file__))
LAST_REG = []
fm = lambda x, n=2: _dec(f"{float(x):.{n}f}")
PB = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'


def tbl(head, rows, al=None):
    al = al or ["c"] * len(head)
    sep = [{"l": ":---", "c": ":---:", "r": "---:"}[a] for a in al]
    out = "| " + " | ".join(_cell(h) for h in head) + " |\n| " + " | ".join(sep) + " |\n"
    for r in rows: out += "| " + " | ".join(str(_cell(x)) for x in r) + " |\n"
    return out + "\n"


def T(text, **kw):
    for k, v in kw.items(): text = text.replace("<<" + k + ">>", str(v))
    return text


def bracket(xs, ys, x):
    """linearna interpolacija – vraća (j, x_j, x_j1, y_j, y_j1, y)"""
    if x <= xs[0]: return None
    if x >= xs[-1]: return None
    j = max(i for i in range(len(xs) - 1) if xs[i] <= x)
    y = ys[j] + (x - xs[j]) / (xs[j + 1] - xs[j]) * (ys[j + 1] - ys[j])
    return j, xs[j], xs[j + 1], ys[j], ys[j + 1], y


def build_report(P, sections, cases, axes, fw, bw, has_caps, F, cfg, outdir, log=print, template_mode=False):
    md = []; A = md.append
    tn = [0]; REG = []
    def newtab(title, meaning=None):
        tn[0] += 1
        if meaning: REG.append((_TF('Tabela {0}', tn[0]), meaning))
        return _TF('**Tabela {0} – {1}**\n\n', tn[0], title)
    Vr, maxVr, a_, d_ = P.Vr, P.maxVr, P.a, P.d
    Rmin_Vr = P.RminVr(); fRVr = P.fRv(Vr)
    proj = cfg.get("projekat", _T('Dionica'))
    chains = [(ch, _T('smjer rastuće stacionaže')) for ch in (fw,) if ch] + [(ch, _T('suprotno smjeru stacionaže')) for ch in (bw,) if ch]
    s_lo = min(c.SA for c in cases); s_hi = max(c.SB for c in cases)
    uniq = []; seen = set()
    for c in cases:
        k = (id(c.axis), c.SA, c.SB)
        if k not in seen: seen.add(k); uniq.append(c)
    def sh(c): return c.shortname
    # --- pomoćne statistike
    def rstat(ch, sel):
        from .engine import combine
        items = [sel(c) for c in ch]; return combine(items)
    rs_car = {id(ch): rstat(ch, lambda c: dict(c.stc, L=c.L)) for ch, _ in chains}
    rs_trk = {(id(ch), q): rstat(ch, lambda c, q=q: dict(c.trk[q]["stat"], L=c.L)) for ch, _ in chains for q in range(len(P.vozila))}
    maxG = {c.key: float(np.abs(c.G).max()) for c in cases}
    # lukovi
    def arcs(c):
        out = []
        for e in c.axis.elems:
            if e["typ"] == "KRUZNI_LUK" and e["s1"] > c.SA and e["s0"] < c.SB and e["L"] > 0.5:
                a = max(e["s0"], c.SA); b = min(e["s1"], c.SB); m = (c.S >= a - 1e-9) & (c.S <= b + 1e-9)
                if not m.any(): continue
                R = abs(e["R0"]); V = float(c.V[m].min()); mid = (a + b) / 2
                RmA = float(np.interp(float(c.VRS[m].min()), P.Vr_tab, P.Rmin_tab))      # Rmin za računsku brzinu na luku (zona Vr ili osnovna)
                ipk = min(7, max(2.5, 7 * (RmA / R) ** 0.74)); ipa = abs(c.axis.cross(mid))
                fR = P.fRv(V); fp = V * V / (127 * R) - ipa / 100; ipr = max(2.5, (V * V / (127 * R) - fR) * 100)
                out.append(dict(no=e["no"], R=R, dir="desno" if e["R0"] > 0 else "lijevo", s0=e["s0"], s1=e["s1"], a=a, b=b, V=V, ipk=ipk, ipa=ipa, fR=fR, fp=fp, ipr=ipr,
                                ok1=ipa >= round(ipk, 1) - 1e-9, ok2=fp <= fR + 1e-9, ok3=R >= RmA - 1e-9, Rmin=RmA, Vrl=float(c.VRS[m].min()), vpR=min(P.vpR(R), float(c.VMX[m].min())), vmx=float(c.VMX[m].min()), partial=(e["s0"] < c.SA - 1 or e["s1"] > c.SB + 1)))
        return out
    ARC = {c.key: arcs(c) for c in cases}
    def unitT(c, un, X):
        a = max(un["s0"], c.SA); b = min(un["s1"], c.SB); m = (c.S >= a - 1e-9) & (c.S <= b + 1e-9)
        return float(X[m].min() if un["kind"] == "K" else X[m].max())
    def ocj(d): return _T("dobro") if d <= 10 else (_T('zadovoljavajuće') if d <= 20 else _T("loše"))
    def unit_rows(c, TT, VV):
        rows = []; un = c.units; res = []
        for i, u in enumerate(un):
            a = max(u["s0"], c.SA); b = min(u["s1"], c.SB); vt = unitT(c, u, TT); vr = unitT(c, u, VV)
            dvt = abs(vt - unitT(c, un[i + 1], TT)) if i < len(un) - 1 else None
            dvr = abs(vr - unitT(c, un[i + 1], VV)) if i < len(un) - 1 else None
            rows.append([i + 1, _T('krivina') if u["kind"] == "K" else "pravac", fm(a, 1), fm(b, 1), fm(b - a, 1), "–" if math.isinf(u["R"]) else fm(u["R"], 2), u.get("dir", "–"),
                         fm(vt, 1), "–" if dvt is None else fm(dvt, 1), "" if dvt is None else ocj(dvt), fm(vr, 1), "–" if dvr is None else fm(dvr, 1)])
            res.append((i, u, vt, vr, dvt, dvr))
        return rows, res
    UN = {c.key: unit_rows(c, c.T, c.V) for c in cases}

    # =========================================================== naslovna
    desc = "; ".join(f"{('st. %s – %s' % (fm(s['od'], 0), fm(s['do'], 3).rstrip('0').rstrip(','))) }: {_T(TIPOVI[s['tip']])}" for s in sections)
    if not template_mode: A(_TF('---\ntoc-title: "Sadržaj"\ntitle: "ELABORAT 3 – VOZNO-DINAMIČKE ANALIZE"\nsubtitle: "{0} – računska brzina V~r~ = {1:.0f} km/h"\nauthor: "Rezultirajući profil projektne brzine, usklađenost elemenata trase, vrijeme putovanja i potrošnja goriva"\ndate: "{2}"\n---\n\n', proj, Vr, cfg.get('datum', '')))
    # =========================================================== 1 Uvod
    A("# Uvod\n")
    A(_TF('U sklopu Glavnog projekta – {0} – za usvojenu trasu urađene su vozno-dinamičke analize u skladu sa Projektnim zadatkom (Elaborat 3). Cilj analize je da se definišu mjerodavni parametri za dimenzionisanje konstruktivnih elemenata plana i profila, da se utvrde eksploatacioni efekti (brzine, vrijeme putovanja, potrošnja goriva) i da se provjeri usklađenost i homogenost primijenjenih elemenata projektne geometrije.\n', proj))
    A(_TF('Analizirani raspon trase je st. {0} – {1} m. Trasa se sastoji od {2}:\n', fm(s_lo, 3), fm(s_hi, 3), _T('sljedećih sekcija') if len(sections) > 1 else _T('jedne sekcije')))
    for sec in sections:
        i = sec["no"]; od, do = sec["od"], sec["do"]; tip = sec["tip"]
        cs = [c for c in cases if c.no == i]
        axn = "; ".join(sorted({f"{c.axis.name}" for c in cs}))
        txt = {"razdvojene": _T('dva odvojena kolovoza, svaki sa svojom osovinom – desni kolovoz (vožnja u smjeru rastuće stacionaže) i lijevi kolovoz (vožnja suprotno smjeru stacionaže)'),
               "centralna": _T('centralna osovina sa jednim zajedničkim ispisom geometrije; proračun je urađen za oba smjera vožnje (u smjeru stacionaže i suprotno)'),
               "jedna_naprijed": _T('jedna osovina, proračun samo za smjer rastuće stacionaže'),
               "jedna_nazad": _T('jedna osovina, proračun samo za smjer opadajuće stacionaže')}[tip]
        zone = [z for c in cs[:1] for z in c.zvr]
        extra = (_T('; zone računske brzine: ') + ", ".join(_TF('V~r~ = {0} km/h na st. {1} – {2}', fm(z[2], 0), fm(z[0], 3), fm(z[1], 3)) for z in zone)) if zone else ""
        A(_TF("- **sekcija {0}: st. {1} – {2} m** – {3} ({4}){5}", i, fm(od, 3), fm(do, 3), txt, axn, extra) + (";" if sec is not sections[-1] else "."))
    A("")
    if cfg.get("opis_sekcija"): A(str(cfg["opis_sekcija"]) + "\n")
    zv = sorted({z for c in cases for z in c.zvr})
    if zv: A(_TF('Računska brzina je promijenjena na zonama: {0}; izvan zona važi V~r~ = {1:.0f} km/h. U zonama se max V~r~ = V~r~ + {2} km/h, R~min~(V~r~), granica V~p,R~ i provjere lukova računaju za računsku brzinu zone, a prijelazi brzine između zona izvode se ubrzanjem a i usporenjem d.\n', "; ".join(_TF('V~r~ = {0} km/h na st. {1} – {2}', fm(z[2], 0), fm(z[0], 3), fm(z[1], 3)) for z in zv), Vr, fm(P.dodatak_maxVr, 0)))
    A(_TF('Sa osnovnom računskom brzinom **V~r~ = {0:.0f} km/h** analiza obuhvata:\n', Vr))
    A(_TF('- rezultirajući profil projektne brzine V~p~ (V85) za putničko vozilo i teška vozila ({0}) u uslovima slobodnog saobraćajnog toka, za sve smjerove vožnje na cijeloj trasi;\n- usklađivanje (harmonizaciju) elemenata trase i statističke pokazatelje dinamičke homogenosti: srednju vrijednost projektne brzine, srednje kvadratno odstupanje i koeficijent dinamičke homogenosti;\n- profil brzina, vrijeme putovanja i potrošnju goriva teških vozila;\n- provjeru poprečnih nagiba kolovoza na kružnim lukovima za dobijene projektne brzine{1}.\n\n**Zahtijevana preglednost i optičke analize (Elaborat 4) nisu predmet ovog elaborata.** Rezultirajući profili projektne brzine iz poglavlja 4 su ulazni podatak za te analize (P~zp~ = f(V~p~)).\n\nNumerički proračun dat je tabelarno u ovom elaboratu (prilozi A) i u pratećoj radnoj knjizi (Excel), u kojoj su svi koraci urađeni formulama, a svi dijagrami iz ovog elaborata izrađeni su kao Excel dijagrami koji se automatski preračunavaju kada se u listove *Ulaz_n* ubace novi ispisi (ulazni podaci su označeni plavom bojom).\n', ', '.join(v['kratko'].lower() for v in P.vozila), ''))
    A(_T("\n# Oznake i simboli {.unnumbered}\n"))
    A(_T('Brojevi u uglastim zagradama [n] u tekstu odnose se na popis literature na kraju elaborata. Indeksi slučajeva (smjerova vožnje): ') + "; ".join(f"**{c.key}** – {c.label}" for c in cases) + _T('.\n\n**Tabela O1 – Značenje oznaka**\n\n'))
    A(tbl(["Oznaka", _T('Značenje'), "Jedinica", "Izvor vrijednosti", "Gdje se koristi"], [list(r) for r in tr_obj(SYM)], ["l", "l", "c", "l", "l"]))
    A("@@TABCOL@@")
    # =========================================================== 2 ulazni podaci
    A(_T("\n# Ulazni podaci i metodologija\n"))
    A(_T("## Ulazni podaci\n"))
    A(_T('Korišteni su ispisi projektanta, po tri datoteke za svaku osovinu (horizontalna osovina, uzdužni profil, poprečni nagibi):\n'))
    A(tbl([_T('Osovina'), "Horizontala", _T('Uzdužni profil'), _T('Poprečni nagibi'), "Br. elem.", _T('Dužina [m]'), "Raspon st. [m]", "Zatvaranje [mm]", _T('K proračun / ispis [gon/km]'), "Nivelete [mm]"],
          [[ax.name, os.path.basename(ax.files[0]), os.path.basename(ax.files[1]), os.path.basename(ax.files[2]), len(ax.elems), fm(ax.length, 3), f"{fm(ax.s_min, 1)} – {fm(ax.s_max, 1)}",
            fm(max(abs(ax.end_err[0]), abs(ax.end_err[1])) * 1000, 0), fm(ax.K, 3) + (" / " + fm(ax.K_file, 3) if ax.K_file else ""), "–" if ax.vdev is None else fm(ax.vdev * 1000, 0)] for ax in axes], ["l"] * 4 + ["c"] * 6))
    A(_TF('Kao provjera ispravnosti preuzimanja podataka, osovine su rekonstruisane numeričkom integracijom elemenata (odstupanje krajnje tačke od ispisa dato je u tabeli), a karakteristika krivine iz ispisa poredi se sa izračunatom. Uzdužni profili su rekonstruisani iz tablica prelomnih tačaka (stacionaža, kota, radijus vertikalne krivine, nagib poslije prelomne tačke, početak i kraj krivine); rekonstruisane kote porede se sa kotama osovine u ispisu poprečnih profila (kolona „Nivelete”).\n\nPoprečni nagib kolovoza za provjeru na krivinama očitava se sa vanjske trake lijeve strane kolovoza; konvencija: pozitivan nagib raste od osovine ka spolja, a nagib desne strane je u ispisima jednak negativnom nagibu lijeve strane (maksimalna razlika |L + D| u ispisima: {0}). Poprečni nagib se koristi samo u provjerama na kružnim lukovima (poglavlje 4.3) i ne ulazi u profil brzine, vrijeme i potrošnju goriva.\n\nProračun je izvršen na mreži Δs = {1} m (u zadnjem koraku sekcije korak može biti kraći).\n', ', '.join(f'{ax.name} {fm(ax.cs_mirror, 3)} %' for ax in axes), fm(P.korak, 0)))
    A(_T('\n## Metoda određivanja projektne brzine V~p~ (V85) – srpske smjernice\n'))
    A(T(_T('Za određivanje brzine V85 primijenjena je metoda iz **Pravilnika o uslovima koje javni putevi izvan naselja moraju da ispunjavaju sa aspekta bezbjednosti saobraćaja (Sl. glasnik RS 50/2011), Prilog 2** [1], uz **Priručnik za projektovanje puteva u Republici Srbiji (Putevi Srbije), dio 4.0** [2] (model vožnje). Prema tim dokumentima projektna brzina V~p~ nije zadana unaprijed, već se dobija kao *posljedica* usvojenih elemenata geometrije, u rasponu\n\n$$V_{ri} \\le V_p \\le \\max V_r, \\qquad \\max V_r = V_r + <<dm>> = <<mx>>\\ \\text{km/h}$$\n\nNajveća brzina na pravcu je max V~r~, odnosno V~ri~ + <<dm>> km/h, pri čemu se ta vrijednost poredi s max V~r~ iz tabele 3-04 Priloga 2 i ne smije je prekoračiti<<tab34>>.\n\nRezultirajući profil V~p~ konstruiše se po stacionaži na sljedeći način:\n\n1. za svaki kružni luk iz zakonitosti **V~p~ – R** određuje se brzina koju vozač ostvaruje u krivini; na pravcima je ciljna brzina max V~r~;\n2. iz zakonitosti **V~p~ – i~N~** (uticaj podužnog nagiba) određuje se ciljna brzina u zavisnosti od nagiba nivelete u smjeru vožnje;\n3. mjerodavna ciljna brzina je manja od te dvije (i od eventualnog ograničenja brzine u scenariju);\n4. ciljna brzina se „glača" modelom vožnje sa ubrzanjem a = <<a>> m/s² i usporenjem d = <<d>> m/s² (nominalno a = d = 0,8 m/s², a put promjene brzine određen je obrascem iz slike 3-09: L = 0,1·V~sr~·ΔV, što odgovara a = d = 0,77 m/s²);\n5. na tako dobijenom profilu određuju se statistički pokazatelji (V~p,sr~, σ, D~h~) i usklađenost susjednih elemenata ΔV~p~.\n\nProfil se konstruiše posebno za svaki smjer vožnje: u smjeru rastuće stacionaže i u smjeru opadajuće stacionaže (podužni nagib i~N~ se uzima sa predznakom u smjeru vožnje, + uzbrdo). Slučajevi (dionice) u istom smjeru vožnje lančano se nastavljaju na granicama sekcija: brzina na kraju jedne dionice je ulazna brzina sljedeće (zadovoljavaju se i ubrzanje i usporenje kroz granicu). Na ulazu u trasu pretpostavljena je slobodna vožnja sa <<ve>> km/h.\n\n**Napomena o zakonitostima.** Zakonitosti V~p~ – R i V~p~ – i~N~ su u Prilogu 2 pravilnika [1] date grafički (slike 3-07 i 3-08): (a) za V~p~ – R – primijenjena je linearna interpolacija tablice minimalnih radijusa R~min~(V~ri~), izračunatih obrascem sa slike 3-07, R~min~ = V~ri~² / (127·(f~R~(V) + i~pk,max~)), sa i~pk,max~ = 7 % i f~R~ iz tabele 3-05, ograničena na max V~r~; (b) za V~p~ – i~N~ – tablica je očitana s dijagrama na slici 3-08 (tačnost očitanja oko ±5 km/h) [4]. Obje zakonitosti su u radnoj knjizi ulazni podaci (list *Parametri*) i mogu se zamijeniti preciznije očitanim vrijednostima bez promjene ostalih formula. Rezultati poglavlja 4 zavise od ovih tablica, pa se preporučuje kontrola očitanih vrijednosti prema originalnim dijagramima.'),
           dm=fm(P.dodatak_maxVr, 0), tab34=(_TF('; za ovaj put max V~r~ prema tabeli 3-04 iznosi {0} km/h, pa je max V~r~ = min(V~r~ + {1}; {2}) = {3} km/h', fm(P.maxVr_tab, 0), fm(P.dodatak_maxVr, 0), fm(P.maxVr_tab, 0), fm(maxVr, 0)) if P.maxVr_tab not in (None, '', 0) else ''), mx=fm(maxVr, 0), a=fm(a_, 2), d=fm(d_, 2), ve=fm(P.ulazna_brzina, 0)))
    A("\n### Zakonitost V~p~ – R\n")
    A(_T('Tabela ') + str(tn[0] + 1) + _T(' sadrži podatke iz Priloga 2 pravilnika: minimalne radijuse horizontalnih krivina i koeficijente trenja za računske brzine V~r~.\n\n'))
    A(newtab(_T('Minimalni radijusi (obrazac sa slike 3-07, i~pk,max~ = 7 %) i koeficijenti trenja (tabela 3-05 Priloga 2 pravilnika [1])'), _T('V~r~ – računska brzina; R~min~ – najmanji radijus; f~T~, f~R~ – dopušteni podužni i bočni koeficijent trenja.')))
    A(tbl(["V~r~ [km/h]"] + [fm(x, 0) for x in P.Vr_tab], [["R~min~ [m]"] + [fm(x, 1) for x in P.Rmin_tab], [_T('f~T~ (podužno)')] + [fm(x) for x in P.fT], [_T('f~R~ (poprečno)')] + [fm(x) for x in P.fR]]))
    A(_T('Brzina za proizvoljni radijus R između tabličnih vrijednosti R~j~ i R~j+1~ dobija se linearnom interpolacijom:\n\n$$V_p(R) = V_j + \\frac{R - R_j}{R_{j+1} - R_j}\\,(V_{j+1} - V_j) \\le \\max V_r$$\n'))
    allarcs = [(c, ar) for c in uniq for ar in ARC[c.key]]
    if allarcs:
        c0, ar0 = min(allarcs, key=lambda x: x[1]["vpR"])
        br = bracket(P.Rmin_tab, P.Vr_tab, ar0["R"])
        if br and br[5] <= maxVr:
            j, R1, R2, V1, V2, y = br
            A(T(_T('**Primjer – kružni luk br. <<no>> ({<<ax>>}, R = <<R>> m, st. <<s0>> – <<s1>>):** R leži između R~j~ = <<R1>> m (V = <<V1>> km/h) i R~j+1~ = <<R2>> m (V = <<V2>> km/h):\n\n$$V_p(<<R>>) = <<V1>> + \\frac{<<R>> - <<R1>>}{<<R2>> - <<R1>>}\\,(<<V2>> - <<V1>>) = <<V1>> + <<q>> \\cdot <<dV>> = <<y>>\\ \\text{km/h}$$\n').replace("{<<ax>>}", "<<ax>>"), no=ar0["no"], ax=c0.axis.name, R=fm(ar0["R"], 1), s0=fm(ar0["s0"], 1), s1=fm(ar0["s1"], 1), R1=fm(R1, 0), R2=fm(R2, 0), V1=fm(V1, 0), V2=fm(V2, 0), q=fm((ar0["R"] - R1) / (R2 - R1), 4), dV=fm(V2 - V1, 0), y=fm(y, 2)))
        else:
            A(_TF('Na svim kružnim lukovima najmanji radijus je R = {0} m (osovina {1}, luk br. {2}), što daje V~p,R~ = {3} km/h (ograničenje max V~r~ = {4} km/h).\n', fm(ar0['R'], 1), c0.axis.name, ar0['no'], fm(ar0['vpR'], 1), fm(maxVr, 0)))
        lim = [(c, ar) for c, ar in allarcs if ar["vpR"] < ar["vmx"] - 1e-6]
        if lim:
            A(_T('Lukovi kod kojih zakonitost V~p~ – R ograničava brzinu: ') + "; ".join(_TF('{0}, luk br. {1} (R = {2} m, st. {3} – {4}) – V~p,R~ = {5} km/h', c.axis.name, ar['no'], fm(ar['R'], 2), fm(ar['s0'], 1), fm(ar['s1'], 1), fm(ar['vpR'], 1)) for c, ar in lim) + _TF('. Svi ostali kružni lukovi imaju V~p,R~ = {0} km/h (ograničenje max V~r~).\n', fm(maxVr, 0)))
        else:
            A(_TF('Nijedan kružni luk ne ograničava brzinu ispod max V~r~ = {0} km/h.\n', fm(maxVr, 0)))
    A("\n### Zakonitost V~p~ – i~N~\n")
    A(newtab(_T('Najveća brzina u zavisnosti od podužnog nagiba (tablica sastavljena na osnovu dijagrama iz Priloga 2 pravilnika, slika 3-08 [4])'), _T('i~N~ – podužni nagib; V~p~(i~N~) – najveća brzina za dati nagib.')))
    A(tbl(["i~N~ [%]"] + [fm(x, 1) for x in P.iN_tab], [["V~p~(i~N~) [km/h]"] + [fm(x, 0) for x in P.ViN_tab]]))
    A(_T('Najveći podužni nagibi nivelete: ') + "; ".join(f"{c.key} ({c.label}) {fm(maxG[c.key], 3)} %" for c in cases) + ".\n")
    gl = [(c, int(np.argmin(c.VPG))) for c in cases if c.VPG.min() < maxVr - 1e-6]
    if gl:
        c, i = min(gl, key=lambda x: x[0].VPG[x[1]]); g = c.sg * c.G[i]; br = bracket(P.iN_tab, P.ViN_tab, g)
        txt = _TF('Zakonitost V~p~ – i~N~ ograničava brzinu u slučajevima: {0}; najnepovoljnije u slučaju {1} ({2}), gdje nagib u smjeru vožnje dostiže {3} %', ', '.join(sorted({x.key for x, _ in gl})), c.key, c.label, fm(g, 3))
        if br:
            j, g1, g2, v1, v2, y = br
            txt += r""":

$$V_p(i_N) = <<v1>> + \frac{<<g>> - <<g1>>}{<<g2>> - <<g1>>}\,(<<v2>> - <<v1>>) = <<y>>\ \text{km/h}$$
""".replace("<<v1>>", fm(v1, 0)).replace("<<v2>>", fm(v2, 0)).replace("<<g>>", fm(g, 3)).replace("<<g1>>", fm(g1, 1)).replace("<<g2>>", fm(g2, 1)).replace("<<y>>", fm(y, 2))
        A(txt + _T('\nZbog toga je potvrda izvora tablice V~p~(i~N~) posebno važna za te dionice.\n'))
    else:
        A(_T('Na analiziranim dionicama zakonitost V~p~ – i~N~ ne ograničava brzinu putničkog vozila (V~p~(i~N~) ≥ max V~r~), pa V~p~ zavisi samo od radijusa lukova.\n'))
    A(_T('\n### Model vožnje\n'))
    A(T(_T('Prijelaz sa jedne ciljne brzine V~1~ na drugu odvija se konstantnim ubrzanjem a ili usporenjem d. Brzina nakon pređenog puta s određuje se iz\n\n$$v_2 = \\sqrt{v_1^2 \\pm 2\\,a\\,s}$$\n\na dužina puta potrebna za promjenu brzine sa V~1~ na V~2~ je\n\n$$s = \\frac{v_1^2 - v_2^2}{2\\,d}$$\n\nPravilnik (slika 3-09) daje put ubrzanja i usporenja pri a = d = 0,8 m/s² obrascem L~a,d~ = 0,1·(V~sr~·ΔV), gdje je V~sr~ = (V~1~ + V~2~)/2 i ΔV = |V~1~ − V~2~| (brzine u km/h, put u m). Taj obrazac je ekvivalentan gornjim jednačinama s ubrzanjem i usporenjem a = d = 1/(0,1·3,6²) = <<aef>> m/s², pa se u proračunu primjenjuju te vrijednosti (npr. 100 → 80 km/h: L = 0,1·90·20 = 180 m).\n'), aef=fm(a_, 4)))
    if allarcs:
        c0, ar0 = min(allarcs, key=lambda x: x[1]["V"]); v2 = ar0["V"]; v1 = maxVr
        if v2 < v1 - 1e-6:
            s_ = ((v1 / 3.6) ** 2 - (v2 / 3.6) ** 2) / (2 * d_)
            A(T(_T('**Primjer – ulaz u kružni luk br. <<no>> (<<ax>>, smanjenje <<v1>> → <<v2>> km/h):**\n\n$$s = \\frac{(<<v1>>/3,6)^2 - (<<v2>>/3,6)^2}{2 \\cdot <<d>>} = \\frac{<<a1>> - <<b1>>}{<<d2>>} = <<s>>\\ \\text{m}$$\n\nPočetak luka je na st. <<s0>> m, pa (u smjeru vožnje koji ide ka luku) usporenje počinje oko <<dist>> m prije početka luka, a nakon izlaza iz luka brzina se vraća na <<v1>> km/h na istom rastojanju.'),
                 no=ar0["no"], ax=c0.axis.name, v1=fm(v1, 0), v2=fm(v2, 1), d=fm(d_, 2), a1=fm((v1 / 3.6) ** 2, 2), b1=fm((v2 / 3.6) ** 2, 2), d2=fm(2 * d_, 3), s=fm(s_, 1), s0=fm(ar0["s0"], 3), dist=fm(s_, 1)))
    A(_T('\n## Statistički pokazatelji dinamičke homogenosti\n'))
    A(_T('Na osnovu rezultirajućeg profila V~p~(s) na dionici dužine L određuju se (ponderisano po dužini, integracija trapeznom formulom na mreži Δs):\n\n$$V_{p,sr} = \\frac{1}{L}\\int_0^L V_p(s)\\,ds \\approx \\frac{1}{L}\\sum_{i=1}^{n-1} \\frac{V_i + V_{i+1}}{2}\\,\\Delta s_i$$\n\n$$\\sigma = \\sqrt{\\frac{1}{L}\\int_0^L \\left(V_p(s) - V_{p,sr}\\right)^2 ds}$$\n\n$$D_h = \\frac{\\sigma}{V_{p,sr}} \\cdot 100\\ [\\%]$$\n\nDinamička usklađenost susjednih elemenata trase (pravac – krivina – pravac) izražava se razlikom projektnih brzina\n\n$$\\Delta V_{i,i+1} = \\left| V_{p,i} - V_{p,i+1} \\right|$$\n\nKriterij ocjene (Lamm i dr. [5]): ΔV ≤ 10 km/h – dobro; 10 < ΔV ≤ 20 km/h – zadovoljavajuće; ΔV > 20 km/h – loše. Za projektnu brzinu elementa uzima se: na kružnom luku (sa prijelaznicama) najmanja brzina, na pravcu najveća brzina u elementu. Ocjena se daje prema razlici *ciljnih* brzina elemenata (V~p,R~ odnosno V~p~(i~N~), prije glačanja ubrzanjem/usporenjem), jer rezultirajući profil po definiciji ublažava skokove; razlika rezultirajućih brzina data je uz nju radi poređenja.\n\nVrijeme putovanja određuje se iz\n\n$$T = \\int_0^L \\frac{ds}{v(s)} \\approx \\sum_{i=1}^{n-1} \\frac{\\Delta s_i}{2}\\left(\\frac{1}{v_i} + \\frac{1}{v_{i+1}}\\right)$$'))
    # =========================================================== 3 geometrija
    A(_T('\n# Geometrija trase\n'))
    A(_T("## Horizontalni elementi\n"))
    def elrows(c):
        rows = []
        for e in c.axis.elems:
            if e["s1"] > c.SA and e["s0"] < c.SB and e["L"] > 0.5:
                typ = {"PRAVAC": "pravac", "KRUZNI_LUK": _T('kružni luk'), "PRIJELAZNICA": _T('prijelaznica')}[e["typ"]]
                Rt = "∞" if e["typ"] == "PRAVAC" else (fm(abs(e["R0"]), 2) if e["typ"] == "KRUZNI_LUK" else f"{'∞' if math.isinf(e['R0']) else fm(abs(e['R0']), 2)} → {'∞' if math.isinf(e['R1']) else fm(abs(e['R1']), 2)}")
                smj = "" if e["typ"] == "PRAVAC" else ("desno" if (e["R0"] if not math.isinf(e["R0"]) else e["R1"]) > 0 else "lijevo")
                rows.append([e["no"], typ, smj, fm(e["s0"], 3), fm(e["s1"], 3), fm(e["L"], 3), Rt, fm(e["A"], 3) if e["A"] else "–"])
        return rows
    HEAD_EL = _T('Br. – redni broj elementa u ispisu; Element – pravac / kružni luk / prijelaznica; Smjer – smjer krivine; St. početka i kraja; L – dužina; R – radijus (za prijelaznicu: od → do, ∞ = pravac); A – parametar klotoide.')
    t_first = tn[0] + 1
    for c in uniq:
        A(newtab(_TF('Elementi horizontalne osovine: {0}, st. {1} – {2} ({3})', c.axis.name, fm(c.SA, 0), fm(c.SB, 3).rstrip('0').rstrip(','), os.path.basename(c.axis.files[0])), None))
        A(tbl(["Br.", "Element", "Smjer", _T('St. početka [m]'), "St. kraja [m]", "L [m]", "R [m]", "A [m]"], elrows(c)))
    REG.append((_TF('Tabele {0}–{1}', t_first, tn[0]) if tn[0] > t_first else _TF('Tabela {0}', t_first), HEAD_EL))
    cl = None
    for c in uniq:
        for e in c.axis.elems:
            if e["typ"] == "PRIJELAZNICA" and e["A"] and (math.isinf(e["R0"]) != math.isinf(e["R1"])) and e["L"] > 1 and e["s1"] > c.SA and e["s0"] < c.SB:
                R = abs(e["R1"] if math.isinf(e["R0"]) else e["R0"]); cl = (c.axis.name, e, R); break
        if cl: break
    if cl:
        n_, e, R = cl
        A(_TF('Kontrola prijelaznica: za klotoidu važi A² = R·L. Npr. {0}, element {1} (R = {2} m, L = {3} m): A = √({4} · {5}) = {6} m, što odgovara ispisu ({7} m).\n', n_, e['no'], fm(R, 2), fm(e['L'], 3), fm(R, 2), fm(e['L'], 3), fm(math.sqrt(R * e['L']), 3), fm(e['A'], 3)))
    part = [(c, e) for c in uniq for e in c.axis.elems if (e["s0"] < c.SA - 0.5 and e["s1"] > c.SA + 0.5) or (e["s0"] < c.SB - 0.5 and e["s1"] > c.SB + 0.5)]
    if part:
        A(_T('Elementi koji prelaze granice sekcija ocjenjuju se samo u dijelu unutar sekcije: ') + "; ".join(_TF("{0}, element {1} (st. {2} – {3}) u sekciji st. {4} – {5}", c.axis.name, e['no'], fm(e['s0'], 1), fm(e['s1'], 1), fm(c.SA, 0), fm(c.SB, 0)) for c, e in part) + _T('. Proračun profila brzine izvršen je na cijeloj dionici (uključujući puni ulaz u luk), a rezultati su prikazani za granice sekcije.\n'))
    low = [(c, ar) for c, ar in allarcs if not ar["ok3"]]
    if low:
        A(_T('**Lukovi sa radijusom manjim od najmanjeg za V~r~ (R~min~(V~r~) = ') + fm(Rmin_Vr, 0) + " m):** " + "; ".join(_TF('{0}, luk br. {1} (R = {2} m, V~p,R~ = {3} km/h)', c.axis.name, ar['no'], fm(ar['R'], 2), fm(ar['vpR'], 1)) for c, ar in low) + _T('. Dinamička provjera poprečnog nagiba data je u poglavlju 4.3.\n'))
    A(_T('\n## Uzdužni profil\n'))
    t_first = tn[0] + 1
    for c in uniq:
        pis = c.axis.pis; rows = []
        for i, p in enumerate(pis[1:-1], 1):
            if not (p[5] > c.SA and p[4] < c.SB): continue
            gin = pis[i - 1][3]; gout = p[3]; dg = gout - gin; L_ = p[5] - p[4]; T_ = L_ / 2; e_ = T_ ** 2 / (2 * p[2])
            rows.append([i, fm(p[0], 3), fm(p[1], 3), fm(gin, 3), fm(gout, 3), fm(abs(dg), 3), fm(p[2], 0), fm(L_, 1), fm(T_, 1), fm(e_, 2), "udubljena" if dg > 0 else _T('ispupčena')])
        A(newtab(_TF('Prelomne tačke nivelete: {0}, st. {1} – {2} ({3})', c.axis.name, fm(c.SA, 0), fm(c.SB, 3).rstrip('0').rstrip(','), os.path.basename(c.axis.files[1]))))
        A(tbl(["PT", "Stac. [m]", "Kota [m]", "i~ul~ [%]", "i~izl~ [%]", "Δi [%]", "R~v~ [m]", "L~vk~ [m]", "T [m]", "e [m]", "Tip"], rows) if rows else _T('(nema vertikalnih krivina u rasponu)\n\n'))
    REG.append((_TF('Tabele {0}–{1}', t_first, tn[0]) if tn[0] > t_first else _TF('Tabela {0}', t_first), _T('PT – prelomna tačka nivelete; Stac., Kota – položaj PT; i~ul~, i~izl~ – nagib prije i poslije; Δi – promjena nagiba; R~v~, L~vk~, T, e – elementi vertikalne krivine; Tip – udubljena (Δi > 0) ili ispupčena (Δi < 0).')))
    ex = None
    for c in uniq:
        pis = c.axis.pis
        for i, p in enumerate(pis[1:-1], 1):
            if p[5] > c.SA and p[4] < c.SB and p[5] > p[4]: ex = (c, i, p, pis[i - 1][3]); break
        if ex: break
    if ex:
        c, i, p, gin = ex; dg = p[3] - gin; Lv = p[5] - p[4]
        z1 = p[1] + gin / 100 * (p[0] - p[4]) * 0 + 0  # kota u PT = kota PT (parabola)
        zpt = c.axis.vert(p[0])[0]
        A(T(_T('Elementi vertikalne krivine (parabola): Δi = i~izl~ − i~ul~; L~vk~ = R~v~·|Δi|/100; T = L~vk~/2; e = T²/(2R~v~). Primjer za PT <<i>> osovine <<ax>> (st. <<s>> m): Δi = <<go>> − (<<gi>>) = <<dg>> %, L~vk~ = <<Rv>> · <<adg>>/100 = <<L1>> m (iz ispisa: <<tk>> − <<tz>> = <<L2>> m), T = <<t>> m, e = <<e>> m.\n\nKota nivelete unutar krivine: z(s) = z~PT~ + i~ul~/100 · (s − s~PT~) + Δi/(200·L~vk~) · (s − s~TZ~)². Nagib nivelete mijenja se linearno unutar vertikalne krivine: i(s) = i~ul~ + Δi·(s − s~TZ~)/L~vk~. Provjera: kota nivelete u tjemenu PT je z~PT~ <<sg>> e = <<zz>> <<sg>> <<e>> = <<zpt>> m.\n'), i=i, ax=c.axis.name, s=fm(p[0], 3), go=fm(p[3], 3), gi=fm(gin, 3), dg=fm(dg, 3), Rv=fm(p[2], 0), adg=fm(abs(dg), 3), L1=fm(p[2] * abs(dg) / 100, 1), tk=fm(p[5], 3), tz=fm(p[4], 3), L2=fm(Lv, 1), t=fm(Lv / 2, 1), e=fm(Lv ** 2 / 4 / (2 * p[2]), 2), zpt=fm(zpt, 3), zz=fm(p[1], 3), sg="+" if dg > 0 else "−"))
    A(_T('Najveći podužni nagib nivelete (apsolutno) po dionicama: ') + "; ".join(f"{c.key}: {fm(maxG[c.key], 3)} %" for c in cases) + ".")
    if max(maxG.values()) > P.max_nagib:
        bad = [(c, (np.abs(c.G) > P.max_nagib).sum() * P.korak) for c in cases if maxG[c.key] > P.max_nagib]
        A(_TF(' Kao referentna granica za V~r~ = {0:.0f} km/h uzima se {1} % (granicu treba potvrditi u pravilniku i Projektnom zadatku); ta granica je premašena na: ', Vr, fm(P.max_nagib, 1)) + "; ".join(f"{c.key} (oko {fm(L_, 0)} m)" for c, L_ in bad) + ".")
    A("\n")
    A(_TF('![]({0}){{width=16cm}}\n\n*Slika 1 – Uzdužni profil nivelete i podužni nagib*\n', os.path.basename(F['uzduzni'])))
    A(_TF('![]({0}){{width=16cm}}\n\n*Slika 2 – Dijagram krivine (1000/R) i poprečnog nagiba lijeve strane*\n', os.path.basename(F['krivina'])))
    A(_TF('![]({0}){{width=14cm}}\n\n*Slika 3 – Situacioni plan osovina sa stacionažama*\n', os.path.basename(F['plan'])))
    # =========================================================== 4 putničko
    A(_T('\n# Rezultirajući profil projektne brzine – putničko vozilo\n'))
    A(_T('## Brzine na elementima i harmonizacija\n'))
    vlow = sorted(((ar["V"], c.key, ar) for c in cases for ar in ARC[c.key] if ar["V"] < maxVr - 0.05), key=lambda x: x[0])
    if vlow:
        A(_TF('Primjenom metode iz poglavlja 2 dobijeni su rezultirajući profili V~p~ na slici 4. Ograničenja brzine ispod max V~r~ = {0} km/h postoje na lukovima: ', fm(maxVr, 0)) + "; ".join(_TF('{0}, luk br. {1} (R = {2} m, st. {3} – {4}) – V~p~ = {5} km/h', k, ar['no'], fm(ar['R'], 1), fm(ar['s0'], 1), fm(ar['s1'], 1), fm(v, 1)) for v, k, ar in vlow) + ".\n")
    else:
        A(_TF('Primjenom metode iz poglavlja 2 dobijeni su rezultirajući profili V~p~ na slici 4. Na svim kružnim lukovima ostvaruje se najveća dopuštena brzina max V~r~ = {0} km/h.\n', fm(maxVr, 0)))
    # provjera donje granice V_ri ≤ V_p (tačka 3.3.3 Priloga 2)
    below = sorted(((ar["V"], c.key, ar) for c in cases for ar in ARC[c.key] if ar["V"] < Vr - 0.05), key=lambda x: x[0])
    bg = [(c, int(np.argmin(c.V))) for c in cases if c.V.min() < Vr - 0.05]
    if below or bg:
        w = []
        if below: w.append(_T('na lukovima ') + "; ".join(_TF('{0}, luk br. {1} (R = {2} m, V~p~ = {3} km/h)', k, ar['no'], fm(ar['R'], 1), fm(v, 1)) for v, k, ar in below))
        if bg: w.append(_T('na dionicama s velikim podužnim nagibom: ') + "; ".join(_TF('{0} (najmanja V~p~ = {1} km/h na st. {2} m)', c.key, fm(float(c.V.min()), 1), fm(float(c.S[i]), 0)) for c, i in bg))
        A(_TF('**Provjera uslova V~ri~ ≤ V~p~.** Prema tački 3.3.3 Priloga 2 pravilnika mora biti V~ri~ ≤ V~p~ ≤ max V~r~. Računska brzina je V~r~ = {0} km/h, a rezultirajuća V~p~ je manja od nje ', fm(Vr, 0)) + " i ".join(w) + _T('. Ti elementi ne zadovoljavaju računsku brzinu i treba ih ispraviti ili obrazložiti; brzina se u profilu ne podiže na V~r~ jer bi to prikrilo stvarno ograničenje.\n'))
    else:
        A(_TF('**Provjera uslova V~ri~ ≤ V~p~.** Na cijeloj trasi je V~p~ ≥ V~r~ = {0} km/h (tačka 3.3.3 Priloga 2 pravilnika).\n', fm(Vr, 0)))
    A(_TF('![]({0}){{width=16cm}}\n\n*Slika 4 – Rezultirajući profili projektne brzine – putničko vozilo', os.path.basename(F['putnicko'])) + "*\n")
    t_first = tn[0] + 1; bad_dv = []
    for c in cases:
        rows, res = UN[c.key]
        A(newtab(_TF('Usklađenost susjednih elemenata, {0}: {1} (putničko vozilo)', c.key, c.label)))
        A(tbl(["Br.", "Element", _T('St. od [m]'), _T('St. do [m]'), "L [m]", "R [m]", "Smjer", "V~p,i~ cilj [km/h]", "ΔV~cilj~ [km/h]", "Ocjena", "V~p,i~ rez. [km/h]", "ΔV~rez~ [km/h]"], rows))
        for i, u, vt, vr, dvt, dvr in res:
            if dvt is not None and dvt > 10: bad_dv.append((c, i, u, dvt, dvr))
    REG.append((_TF('Tabele {0}–{1}', t_first, tn[0]) if tn[0] > t_first else _TF('Tabela {0}', t_first), _T('Br., Element – krivina (luk sa prijelaznicama) ili pravac; St. od / do; L; R; Smjer; V~p,i~ cilj – ciljna brzina elementa T (krivina: najmanja, pravac: najveća); ΔV~cilj~ – razlika u odnosu na sljedeći element i Ocjena prema kriteriju 10 / 20 km/h; V~p,i~ rez. i ΔV~rez~ – isto za rezultirajući profil poslije ubrzanja/usporenja.')))
    if not bad_dv:
        A(_T('Na svim dionicama razlike ciljnih brzina susjednih elemenata su ≤ 10 km/h, tj. trasa je sa aspekta dinamičke usklađenosti susjednih elemenata ocijenjena kao **dobra**.\n'))
    else:
        A(_T('Razlike ciljnih brzina susjednih elemenata veće od 10 km/h: ') + "; ".join(_TF('{0}, prelaz sa elementa {1} na element {2} (st. {3}): ΔV~cilj~ = {4} km/h (**{5}**), ΔV~rez~ = {6} km/h', c.key, i + 1, i + 2, fm(u['s1'], 1), fm(d, 1), ocj(d), fm(r, 1)) for c, i, u, d, r in bad_dv) + _T('. Razlika rezultirajućih brzina je manja jer se usporenje odvija na prijelaznici i prethodnim elementima; ostale razlike su ≤ 10 km/h („dobro”).\n'))
    A(_T('\n## Statistika rezultirajućeg profila\n'))
    c = cases[0]; s_ = c.stc
    A(T(_T('Primjer proračuna statistike za slučaj **<<k>>** (<<lab>>), dužina L = <<L>> m, <<n>> tačaka:\n\n$$V_{p,sr} = \\frac{1}{<<L>>}\\sum \\frac{V_i + V_{i+1}}{2}\\,\\Delta s_i = <<m>>\\ \\text{km/h}$$\n\n$$\\sigma = \\sqrt{\\frac{1}{<<L>>}\\sum \\frac{(V_i - <<m>>)^2 + (V_{i+1} - <<m>>)^2}{2}\\,\\Delta s_i} = <<sd>>\\ \\text{km/h}$$\n\n$$D_h = \\frac{\\sigma}{V_{p,sr}} \\cdot 100 = \\frac{<<sd>>}{<<m>>} \\cdot 100 = <<dh>>\\ \\%$$\n\n$$T = \\sum \\frac{\\Delta s_i}{2}\\left(\\frac{1}{v_i} + \\frac{1}{v_{i+1}}\\right) = <<t>>\\ \\text{s} = <<tm>>\\ \\text{min}$$\n'), k=c.key, lab=c.label, L=fm(c.L, 3).rstrip("0").rstrip(","), n=c.N, m=fm(s_["mean"], 3), sd=fm(s_["sd"], 3), dh=fm(s_["Dh"], 2), t=fm(s_["time"], 1), tm=fm(s_["time"] / 60, 2)))
    A(newtab(_T('Pokazatelji dinamičke homogenosti, putničko vozilo'), _T('Dionica / smjer; V~p,sr~, σ, D~h~ – statistika po dužini; V~min~, V~max~; T – vrijeme putovanja; redovi „cijela trasa” obuhvataju sve dionice u istom smjeru vožnje (lančano).')))
    rows = []
    for c in cases:
        x = c.stc; rows.append([f"{c.key}: {c.shortname} ({fm(c.SA, 0)}–{fm(c.SB, 0)})", fm(x["mean"], 2), fm(x["sd"], 3), fm(x["Dh"], 2), fm(x["vmin"], 1), fm(x["vmax"], 1), fm(x["time"], 1), fm(x["time"] / 60, 2)])
    for ch, lab in chains:
        if len(ch) > 1:
            x = rs_car[id(ch)]; rows.append([_TF('**Cijela trasa – {0} ({1})**', lab, run_label(ch) if False else '+'.join(c.key for c in ch)), fm(x["mean"], 2), fm(x["sd"], 3), fm(x["Dh"], 2), fm(x["vmin"], 1), fm(x["vmax"], 1), fm(x["time"], 1), fm(x["time"] / 60, 2)])
    A(tbl([_T('Dionica (smjer vožnje)'), "V~p,sr~ [km/h]", "σ [km/h]", "D~h~ [%]", "V~min~ [km/h]", "V~max~ [km/h]", "T [s]", "T [min]"], rows))
    dhs = [(c.key, c.stc["Dh"]) for c in cases]
    A(_T('Koeficijent dinamičke homogenosti D~h~ po dionicama: ') + "; ".join(f"{k} {fm(v, 2)} %" for k, v in dhs) + _T('. Vrijednosti D~h~ su direktno posljedica pretpostavljenih zakonitosti V~p~–R i V~p~–i~N~ (poglavlje 2.2): ako se iz originalnih grafika očitaju niže brzine u krivinama, profil će biti niži i raznolikiji (veći D~h~).\n'))
    A(_T('\n## Provjera poprečnih nagiba na kružnim lukovima\n'))
    A(T(_T('Potreban poprečni nagib na krivini za računsku brzinu određuje se iz Priloga 2 [1]:\n\n$$i_{pk}(V_r) = 7 \\cdot \\left(\\frac{R_{min}(V_r)}{R}\\right)^{0,74},\\qquad 2,5\\ \\% \\le i_{pk} \\le 7\\ \\%$$\n\na dinamička provjera za rezultirajuću brzinu V (uslov bočnog trenja):\n\n$$f_{pot} = \\frac{V^2}{127\\,R} - \\frac{i_p}{100} \\le f_R(V),\\qquad i_{p,potr} = \\left(\\frac{V^2}{127\\,R} - f_R(V)\\right)\\cdot 100$$\n')))
    if allarcs:
        c0, ar = min(allarcs, key=lambda x: x[1]["fR"] - x[1]["fp"])
        A(T(_T('**Primjer – najnepovoljniji luk: <<k>>, luk br. <<no>> (R = <<R>> m, V = <<V>> km/h, R~min~(<<Vr>>) = <<Rm>> m, primijenjeni nagib <<ip>> %):**\n\n$$i_{pk}(<<Vr>>) = 7 \\cdot \\left(\\frac{<<Rm>>}{<<R>>}\\right)^{0,74} = 7 \\cdot <<x>> = <<y>>\\ \\% <<cmp>>$$\n\n$$f_{pot} = \\frac{<<V>>^2}{127 \\cdot <<R>>} - <<ipf>> = <<a>> - <<ipf>> = <<b>> <<c2>> f_R(<<V>>) = <<fr>>$$\n\n$$i_{p,potr} = \\left(<<a>> - <<fr>>\\right) \\cdot 100 = <<ipr>>\\ \\%$$\n'), k=c0.key, no=ar["no"], R=fm(ar["R"], 2), V=fm(ar["V"], 1), Vr=fm(Vr, 0), Rm=fm(Rmin_Vr, 0), ip=fm(ar["ipa"], 1), x=fm((Rmin_Vr / ar["R"]) ** 0.74, 3), y=fm(7 * (Rmin_Vr / ar["R"]) ** 0.74, 3),
                 cmp=(r"\Rightarrow\ " + fm(ar["ipa"], 1) + r"\ \%\ (\text{primijenjeno})") if 7 * (Rmin_Vr / ar["R"]) ** 0.74 <= 7 else _T('> 7,0\\ \\%\\ (\\text{ne može se postići ni sa najvećim dopuštenim nagibom})'),
                 ipf=fm(ar["ipa"] / 100, 3), a=fm(ar["V"] ** 2 / (127 * ar["R"]), 4), b=fm(ar["fp"], 3), c2=r"\le" if ar["ok2"] else ">", fr=fm(ar["fR"], 3), ipr=fm(ar["ipr"], 2)))
    t_first = tn[0] + 1
    for c in cases:
        A(newtab(_TF('Poprečni nagibi na kružnim lukovima, {0}: {1}', c.key, c.label)))
        A(tbl([_T('Luk br.'), "R [m]", "Smjer", "V~p~ [km/h]", "i~pk~(V~r~) [%]", "i~p~ prim. [%]", "i~p~ ≥ i~pk~", "f~R~(V~p~)", "i~p,potr~ [%]", "f~pot~", "f~pot~ ≤ f~R~", "R ≥ R~min~"],
                  [[ar["no"], fm(ar["R"], 2), ar["dir"], fm(ar["V"], 1), fm(ar["ipk"], 2), fm(ar["ipa"], 1), "da" if ar["ok1"] else "NE", fm(ar["fR"], 3), fm(ar["ipr"], 2), fm(ar["fp"], 3), "da" if ar["ok2"] else "NE", "da" if ar["ok3"] else "NE"] for ar in ARC[c.key]]) if ARC[c.key] else _T('(nema kružnih lukova)\n\n'))
    REG.append((_TF('Tabele {0}–{1}', t_first, tn[0]) if tn[0] > t_first else _TF('Tabela {0}', t_first), _T('Luk br. – broj elementa; R; Smjer; V~p~ – najmanja brzina na luku (u dijelu unutar dionice); i~pk~(V~r~) – potreban nagib za V~r~; i~p~ prim. – primijenjeni nagib; i~p~ ≥ i~pk~ – uslov (poređenje na 0,1 %); f~R~(V~p~); i~p,potr~ – potreban nagib za V~p~; f~pot~ – iskorišteno trenje; f~pot~ ≤ f~R~ – uslov; R ≥ R~min~ – da li je radijus veći od najmanjeg za V~r~.')))
    bad1 = [(c, ar) for c in cases for ar in ARC[c.key] if not ar["ok1"]]
    bad2 = [(c, ar) for c in cases for ar in ARC[c.key] if not ar["ok2"]]
    bad3 = [(c, ar) for c in cases for ar in ARC[c.key] if not ar["ok3"]]
    q = lambda l, f: "; ".join(f(c, ar) for c, ar in l)
    out = [_T('**Zaključak provjere:**')]
    out.append(_T('kriterij i~p~ ≥ i~pk~(V~r~) zadovoljen je na svim lukovima.') if not bad1 else _T('kriterij i~p~ ≥ i~pk~(V~r~) nije zadovoljen na: ') + q(bad1, lambda c, ar: _TF('{0}, luk br. {1} (primijenjeno {2} %, potrebno {3} %)', c.key, ar['no'], fm(ar['ipa'], 1), fm(ar['ipk'], 2))) + ".")
    out.append(_T('Dinamički uslov bočnog trenja zadovoljen je na svim lukovima za dobijene brzine.') if not bad2 else _T('Dinamička provjera za rezultirajuću brzinu V~p~ pokazuje da nagib nije dovoljan na: ') + q(bad2, lambda c, ar: _TF('{0}, luk br. {1} (R = {2} m, V = {3} km/h, primijenjeno {4} %, potrebno {5} %)', c.key, ar['no'], fm(ar['R'], 2), fm(ar['V'], 1), fm(ar['ipa'], 1), fm(ar['ipr'], 2))) + _T('. Preporuka: nagib na luku povećati (do 7,0 %), ili prihvatiti manju brzinu na luku, ili povećati radijus.'))
    out.append(_T('Svi lukovi imaju R ≥ R~min~(V~r~).') if not bad3 else _TF('Uslov R ≥ R~min~(V~r~) = {0} m nije zadovoljen na: ', fm(Rmin_Vr, 0)) + q(bad3, lambda c, ar: _TF('{0}, luk br. {1} (R = {2} m)', c.key, ar['no'], fm(ar['R'], 2))) + _T('. Projektant treba potvrditi da je za te lukove dopušteno odstupanje ili povećati radijus.'))
    A(" ".join(out) + "\n")
    # =========================================================== 5 teška
    A(_T('\n# Dinamički model vozila – brzine, vrijeme putovanja i potrošnja goriva\n'))
    A(_T('## Model kretanja i potrošnje\n'))
    vt = "; ".join(f"{v['naziv']}: m = {fm(v['m'] / 1000, 2)} t, P = {fm(v['P'], 0)} kW, C~d~A = {fm(v['CdA'], 2)} m², v~lim~ = {fm(v['vlim'], 0)} km/h, η = {fm(v['eta'], 2)}, γ~m~ = {fm(v['gam'], 2)}, μ = {fm(v['mu'], 2)}, b = {fm(v['b'], 2)} l/kWh, prazan hod {fm(v['prazan_hod'], 1)} l/h" for v in P.vozila)
    v0 = P.vozila[0]
    A(_TF('Za putničko vozilo i teška vozila nisu zadati parametri u Projektnom zadatku, pa su usvojeni tipični (potrebno ih je potvrditi): {0}; gustina vazduha ρ = {1} kg/m³. Profil brzine putničkog vozila iz poglavlja 4 (V~p~) je rezultat zakonitosti V~p~ – R i V~p~ – i~N~ s ograničenjima ubrzanja i usporenja, dok se u ovom poglavlju isto vozilo analizira dinamičkim modelom (sile otpora i vuče), pa se dobijaju brzina po sili vuče, vrijeme putovanja i potrošnja goriva. Na ulazu u trasu vozilo ima v~lim~, a na granicama sekcija brzina se prenosi s prethodne dionice u smjeru vožnje.\n\nZa svaki korak Δs, prema [8] i [9] (v u m/s, G = podužni nagib kao razlomak, + uzbrdo u smjeru vožnje):\n', vt, fm(P.rho, 3)))
    A(_T('$$R_{r} = f_{rl}\\, m\\, g,\\quad f_{rl} = 0,01\\left(1 + \\frac{v}{44,73}\\right)$$\n\n$$R_{a} = \\tfrac{1}{2}\\,\\rho\\, C_d A\\, v^2,\\qquad R_{g} = m\\, g\\, G,\\qquad F_{r} = R_r + R_a + R_g$$\n\n$$F_t = \\min\\left(\\frac{\\eta\\,P}{v},\\ \\mu\\, m\\, g\\right),\\qquad a = \\frac{F_t - F_r}{\\gamma_m\\, m},\\qquad v_{i+1} = \\sqrt{v_i^2 + 2\\,a\\,\\Delta s} \\le v_{lim}$$\n\nPotrošnja goriva po koraku: za F~r~ ≤ 0 (kočenje motorom, bez ubrizgavanja) računa se potrošnja praznog hoda; u protivnom\n\n$$\\Delta B = b \\cdot \\frac{F_u\\, v_{sr}}{\\eta \\cdot 1000} \\cdot \\frac{\\Delta s}{v_{sr}} \\cdot \\frac{1}{3600}\\ [\\text{l}]$$\n\npri čemu je F~u~ = F~r~ kada vozilo održava v~lim~, a F~u~ = F~t~ kada vozilo koristi punu snagu (uspon sa padom brzine ili ubrzavanje).\n'))
    kv = [v for v in P.vozila if v.get("Ft_kriva")]
    if kv:
        A(_TF('**Vozilo prema Pravilniku.** Za merodavno teretno vozilo Pravilnik (Prilog 2, slika 3-10) daje V-L dijagram brzine na usponima, iz kojeg proizlaze ravnotežne brzine od 68, 55, 43, 33, 24 i 15 km/h na usponima od 2 % do 7 %. Za vozilo „{0}” sila vuče nije zadata konstantnom snagom nego krivom F~t~(v), dobijenom iz tih ravnotežnih tačaka (pri ravnoteži je F~t~ = F~r~) za masu m = {1} t i C~d~A = {2} m². Ispod najmanje brzine kriva se zadržava na zadnjoj vrijednosti, a iznad najveće vrijedi konstantna snaga. Time se brzine tog vozila na usponima poklapaju s dijagramom iz Pravilnika. Ostala vozila („{3}”) računaju se s konstantnom snagom (savremena vozila), pa su njihove brzine na usponima veće; njihovi parametri su pretpostavke.\n', kv[0]['kratko'], fm(kv[0]['m'] / 1000, 1), fm(kv[0]['CdA'], 1), ', '.join(v['kratko'] for v in P.vozila if not v.get('Ft_kriva'))))
        A(newtab(_T('Kriva sile vuče vozila prema Pravilniku'), _T('v – brzina; F~t~ – sila vuče izvedena iz ravnotežnih tačaka dijagrama (slika 3-10).')))
        A(tbl(["v [km/h]"] + [fm(p[0], 0) for p in kv[0]["Ft_kriva"]], [["F~t~ [N]"] + [fm(p[1], 0) for p in kv[0]["Ft_kriva"]]]))
        def _eq(veh, gr):
            W = veh["m"] * P.g; lo, hi = 0.5, 45.0
            f = lambda v: engine.ft_of(veh, v, W) - (P.frl0 * (1 + v / P.frl_v) * W + 0.5 * P.rho * veh["CdA"] * v * v + W * gr / 100)
            if f(lo) < 0: return None
            if f(hi) > 0: return hi * 3.6
            for _ in range(60):
                mid = (lo + hi) / 2
                if f(mid) > 0: lo = mid
                else: hi = mid
            return lo * 3.6
        gs = [2, 3, 4, 5, 6, 7]
        A(newtab(_T('Ravnotežne brzine vozila na usponu (model) i dijagram iz Pravilnika'), _T('Brzina pri kojoj sila vuče izjednačava sve otpore; „Pravilnik” – vrijednosti sa slike 3-10.')))
        A(tbl([_T('Uspon [%]')] + [str(g) for g in gs], [[_T('Pravilnik, slika 3-10 [km/h]')] + [str(x) for x in (68, 55, 43, 33, 24, 15)]] + [[f"{v['kratko']} [km/h]"] + [("–" if _eq(v, g) is None else fm(_eq(v, g), 0)) for g in gs] for v in P.vozila]))
    def tex(c, st_, k):
        S = c.S; sgn = c.sg; si = int(np.argmin(abs(S - st_))); veh = P.vozila[k]; W = veh["m"] * P.g
        si = min(max(si, 0), c.N - 1)
        nx = si + sgn
        if nx < 0 or nx >= c.N: si = si - sgn; nx = si + sgn
        v = c.trk[k]["v"][si] / 3.6; ds = abs(S[nx] - S[si]); Gm = sgn * (c.G[si] + c.G[nx]) / 200
        frl = P.frl0 * (1 + v / P.frl_v); Rr = frl * W; Ra = 0.5 * P.rho * veh["CdA"] * v ** 2; Rg = W * Gm; Fr = Rr + Ra + Rg
        Ft = engine.ft_of(veh, v, W); vl = veh["vlim"] / 3.6
        full = not (v >= vl - 1e-9 and Ft >= Fr)
        a2 = (Ft - Fr) / (veh["gam"] * veh["m"]); vn = min(vl, math.sqrt(max(v * v + 2 * a2 * ds, 9))) if full else vl
        vav = (v + vn) / 2; dt = ds / vav; Fu = (Ft if Fr > 0 else 0) if full else max(Fr, 0)
        fl = veh["prazan_hod"] * dt / 3600 if Fr <= 0 else veh["b"] * (Fu * vav / veh["eta"] / 1000) * dt / 3600
        return dict(s=float(S[si]), ds=ds, v=v, Gm=Gm, frl=frl, Rr=Rr, Ra=Ra, Rg=Rg, Fr=Fr, Ft=Ft, a=a2, vn=vn, vav=vav, dt=dt, fl=fl, Fu=Fu, veh=veh, W=W, full=full)
    def exm(E, title, c):
        m = int(E["veh"]["m"])
        return T(_T('**<<title>>** st. <<s>> m: brzina v = <<vkmh>> km/h = <<v>> m/s, nagib u smjeru vožnje G = <<G>> %, korak Δs = <<ds>> m.\n\n$$f_{rl} = 0,01\\left(1 + \\frac{<<v>>}{44,73}\\right) = <<frl>>,\\qquad R_r = <<frl>> \\cdot <<m>> \\cdot <<g>> = <<Rr>>\\ \\text{N}$$\n\n$$R_a = 0,5 \\cdot <<rho>> \\cdot <<cda>> \\cdot <<v>>^2 = <<Ra>>\\ \\text{N},\\qquad R_g = <<m>> \\cdot <<g>> \\cdot <<Gf>> = <<Rg>>\\ \\text{N}$$\n\n$$F_r = <<Rr>> + <<Ra>> + (<<Rg>>) = <<Fr>>\\ \\text{N}$$\n\n$$F_t = \\min\\left(\\frac{<<eta>> \\cdot <<Pw>>}{<<v>>},\\ <<mu>> \\cdot <<m>> \\cdot <<g>>\\right) = \\min(<<Fp>>;\\ <<Fm>>) = <<Ft>>\\ \\text{N}$$\n\n<<tail>>\n'), title=title, s=fm(E["s"], 0), vkmh=fm(E["v"] * 3.6, 1), v=fm(E["v"], 3), G=fm(E["Gm"] * 100, 3), ds=fm(E["ds"], 0), frl=fm(E["frl"], 5), m=m, g=fm(P.g, 2), rho=fm(P.rho, 3),
                 Rr=fm(E["Rr"], 0), cda=fm(E["veh"]["CdA"], 1), Ra=fm(E["Ra"], 0), Gf=fm(E["Gm"], 5), Rg=fm(E["Rg"], 0), Fr=fm(E["Fr"], 0), eta=fm(E["veh"]["eta"], 2), Pw=fm(E["veh"]["P"] * 1000, 0),
                 mu=fm(E["veh"]["mu"], 2), Fp=fm(E["veh"]["eta"] * E["veh"]["P"] * 1000 / E["v"], 0), Fm=fm(E["veh"]["mu"] * E["veh"]["m"] * P.g, 0), Ft=fm(E["Ft"], 0),
                 tail=(T(r"""$$a = \frac{<<Ft>> - (<<Fr>>)}{<<gam>> \cdot <<m>>} = <<a>>\ \text{m/s}^2,\qquad v_{i+1} = \sqrt{<<v>>^2 + 2 \cdot (<<a>>) \cdot <<ds>>} = <<vn>>\ \text{m/s}$$

$$\Delta B = <<b>> \cdot \frac{<<Fu>> \cdot <<vav>>}{<<eta>> \cdot 1000} \cdot \frac{<<dt>>}{3600} = <<fl>>\ \text{l}$$""", Ft=fm(E["Ft"], 0), Fr=fm(E["Fr"], 0), gam=fm(E["veh"]["gam"], 2), m=m, a=fm(E["a"], 4), v=fm(E["v"], 3), ds=fm(E["ds"], 0), vn=fm(E["vn"], 3), b=fm(E["veh"]["b"], 2), Fu=fm(E["Fu"], 0), vav=fm(E["vav"], 3), eta=fm(E["veh"]["eta"], 2), dt=fm(E["dt"], 3), fl=fm(E["fl"], 5))
                       if E["full"] else T(_T('Vozilo je na ograničenju brzine v~lim~ i sile pogona dovoljno: F~u~ = F~r~ = <<Fu>> N, v~i+1~ = v~lim~ = <<vn>> m/s, ΔB = <<fl>> l'), Fu=fm(E["Fu"], 0), vn=fm(E["vn"], 3), fl=fm(E["fl"], 5)) if E["Fr"] > 0 else
                       T(_T('F~r~ ≤ 0 – vozilo ne troši gorivo za pogon (kočenje motorom), računa se samo prazan hod: ΔB = <<b>> · <<dt>>/3600 = <<fl>> l'), b=fm(E["veh"]["prazan_hod"], 1), dt=fm(E["dt"], 3), fl=fm(E["fl"], 5))))
    # primjeri: najslabiji uspon (min brzina) i nizbrdica (Fr<=0) ili puna snaga
    k0 = 0
    cmin = min(cases, key=lambda c: c.trk[k0]["v"].min()); imin = int(np.argmin(cmin.trk[k0]["v"]))
    A(exm(tex(cmin, cmin.S[imin], k0), f"Primjer 1 – {P.vozila[k0]['kratko'].lower()}, {c_ and ''}" if False else _TF('Primjer 1 – {0}, {1} ({2}, najmanja brzina)', P.vozila[k0]['kratko'].lower(), cmin.key, cmin.label.split(',')[0].lower()), cmin))
    done = False
    for c in cases:
        for i in range(0, c.N - 1, 5):
            E = tex(c, c.S[i], k0)
            if E["Fr"] <= 0:
                A(exm(E, f"Primjer 2 – {P.vozila[k0]['kratko'].lower()}, {c.key} (nizbrdica)", c)); done = True; break
        if done: break
    A(_T("\n## Rezultati\n"))
    # tekst
    mins = [(c.trk[k0]["v"].min(), c) for c in cases]
    A(_TF('Najmanje brzine vozila „{0}” po dionicama: ', P.vozila[k0]['kratko']) + "; ".join(f"{c.key} {fm(v, 1)} km/h" for v, c in mins) + _TF('. Na dionicama gdje je vozilo pretežno nizbrdo brzina je ograničena na v~lim~ = {0} km/h; na usponima brzina pada do ravnotežne vrijednosti, što znači veliku razliku prema putničkim vozilima. Potrebu za dodatnom trakom za spora vozila treba ocijeniti prema važećem pravilniku.\n', fm(v0['vlim'], 0)))
    A(_TF('![]({0}){{width=16cm}}\n\n*Slika 5 – Profil brzina teških vozila (siva linija – putničko vozilo)*\n', os.path.basename(F['teska'])))
    A(_TF('![]({0}){{width=16cm}}\n\n*Slika 6 – Kumulativno vrijeme putovanja i potrošnja goriva po pređenom putu*\n', os.path.basename(F['vrijeme'])))
    A(newtab(_T('Vrijeme putovanja i potrošnja goriva teških vozila'), _T('V~sr~, σ, D~h~, V~min~ – statistika brzine teškog vozila; T – vrijeme; Gorivo – ukupno [l], po km i na 100 km; redovi „cijela trasa” su zbir dionica u smjeru vožnje.')))
    rows = []
    for c in cases:
        for q, v in enumerate(P.vozila):
            x = c.trk[q]["stat"]; r = c.trk[q]; Lk = c.L / 1000
            rows.append([f"{v['kratko']} – {c.key}: {c.shortname}", fm(x["mean"], 1), fm(x["sd"], 2), fm(x["Dh"], 2), fm(x["vmin"], 1), fm(r["T"], 1), fm(r["T"] / 60, 2), fm(r["Fu"], 2), fm(r["Fu"] / Lk, 3), fm(r["Fu"] / Lk * 100, 1)])
    for ch, lab in chains:
        if len(ch) > 1:
            for q, v in enumerate(P.vozila):
                x = rs_trk[(id(ch), q)]; Lk = x["L"] / 1000
                rows.append([_TF('**{0} – cijela trasa, {1}**', v['kratko'], lab), fm(x["mean"], 1), fm(x["sd"], 2), fm(x["Dh"], 2), fm(x["vmin"], 1), fm(x["time"], 1), fm(x["time"] / 60, 2), fm(x["fuel"], 2), fm(x["fuel"] / Lk, 3), fm(x["fuel"] / Lk * 100, 1)])
    A(tbl([_T('Vozilo / dionica'), "V~sr~ [km/h]", "σ [km/h]", "D~h~ [%]", "V~min~ [km/h]", "T [s]", "T [min]", "Gorivo [l]", "[l/km]", "[l/100 km]"], rows))
    tot = []
    for q, v in enumerate(P.vozila):
        tt = sum(rs_trk[(id(ch), q)]["time"] for ch, _ in chains); ff = sum(rs_trk[(id(ch), q)]["fuel"] for ch, _ in chains)
        tot.append(f"{v['kratko'].lower()} {fm(tt / 60, 2)} min i {fm(ff, 2)} l")
    A(_T('Ukupno za obilazak cijele trase (svi analizirani smjerovi): ') + "; ".join(tot) + _T('. Potrošnja u smjeru nizbrdice je mala jer se vozilo kreće uz kočenje motorom, a potrošnju u smjeru uspona određuje rad motora punom snagom. Apsolutne vrijednosti zavise od usvojenih parametara vozila (poglavlje 5.1) i služe za poređenje varijanti, a ne kao garantovana potrošnja.\n'))
    # =========================================================== zaključci
    A(_T('\n# Zaključci i preporuke\n'))
    z = []
    for c in cases:
        x = c.stc
        z.append(f"**{c.key} ({c.label}):** V~p~ = {fm(x['vmin'], 1)} – {fm(x['vmax'], 1)} km/h, V~p,sr~ = {fm(x['mean'], 1)} km/h, σ = {fm(x['sd'], 2)} km/h, D~h~ = {fm(x['Dh'], 2)} %, vrijeme putovanja {fm(x['time'] / 60, 2)} min.")
    zk = [_T('Rezultirajuća projektna brzina (V85) i statistika po dionicama: ') + " ".join(z)]
    if any(x[1]["V"] > Vr + 0.05 for x in [(c, ar) for c in cases for ar in ARC[c.key]]) or any(c.V.max() > Vr for c in cases):
        zk.append(_TF('Na dionicama sa povoljnim elementima dobijena V~p~ prelazi računsku brzinu V~r~ = {0} km/h (do {1} km/h). Preglednost u Elaboratu 4 treba provjeriti za V~p~, a ne samo za V~r~.', fm(Vr, 0), fm(maxVr, 0)))
    if bad_dv: zk.append(_T('Dinamička usklađenost susjednih elemenata nije dobra na: ') + "; ".join(_TF("{0} (elementi {1}–{2}, ΔV~cilj~ = {3} km/h – {4})", c.key, i + 1, i + 2, fm(d, 1), ocj(d)) for c, i, u, d, r in bad_dv) + _T('. Preporučuje se ujednačiti radijuse susjednih elemenata ili obezbijediti odgovarajuću signalizaciju.'))
    else: zk.append(_T('Dinamička usklađenost susjednih elemenata je ocijenjena kao dobra (ΔV~cilj~ ≤ 10 km/h) na svim dionicama.'))
    if bad3: zk.append(_TF('Radijus manji od R~min~(V~r~) = {0} m imaju: ', fm(Rmin_Vr, 0)) + "; ".join(_TF('{0}, luk br. {1} (R = {2} m)', c.key, ar['no'], fm(ar['R'], 2)) for c, ar in bad3) + _T('. Preporučuje se povećati radijus ili obrazložiti odstupanje (preporučena brzina, signalizacija).'))
    if bad2 or bad1: zk.append(_T('Poprečni nagib nije dovoljan na: ') + "; ".join(_TF('{0}, luk br. {1} (primijenjeno {2} %, potrebno {3} % za V = {4} km/h)', c.key, ar['no'], fm(ar['ipa'], 1), fm(ar['ipr'], 2), fm(ar['V'], 1)) for c, ar in (bad2 or bad1)) + _T('. Preporuka: povećati nagib (do 7,0 %) ili prihvatiti manju brzinu na luku.'))
    else: zk.append(_T('Primijenjeni poprečni nagibi zadovoljavaju dinamički uslov bočnog trenja na svim lukovima za dobijene brzine.'))
    if max(maxG.values()) > P.max_nagib: zk.append(_TF('Podužni nagib dostiže {0} %, što je više od referentne granice {1} % za V~r~ = {2} km/h (granicu potvrditi u pravilniku/Projektnom zadatku).', fm(max(maxG.values()), 3), fm(P.max_nagib, 1), fm(Vr, 0)))
    if gl: zk.append(_T('Zakonitost V~p~ – i~N~ određuje brzinu na dijelu trase (uspon), pa izvor te zakonitosti treba potvrditi.'))
    tt = []
    for ch, lab in chains:
        tt.append(_TF('{0}: putničko {1} min; ', lab, fm(rs_car[id(ch)]['time'] / 60, 2)) + "; ".join(f"{v['kratko'].lower()} {fm(rs_trk[(id(ch), q)]['time'] / 60, 2)} min / {fm(rs_trk[(id(ch), q)]['fuel'], 2)} l" for q, v in enumerate(P.vozila)))
    zk.append(_T('Vrijeme putovanja i potrošnja goriva po smjerovima: ') + " | ".join(tt) + ".")
    zk.append(_TF('Brzine teških vozila u smjeru uspona padaju do {0} km/h; razlika prema putničkim vozilima jedan je od ulaznih podataka za analizu preticanja, dodatne trake za spora vozila i saobraćajne opreme.', fm(min(c.trk[0]['v'].min() for c in cases), 0)))
    for i, t in enumerate(zk, 1): A(f"{i}. {t}")
    A(_T('\n# Sadržaj i uslovi obrade\n'))
    A(_T('Elaborat je urađen u skladu sa Projektnim zadatkom (Elaborat vozno-dinamičkih analiza) i uslovima obrade iz Opšteg dijela. Numerički proračun dat je u tabelama (prilozi A) i u radnoj knjizi (Excel) sa formulama, dijagramima i svim tabelama iz ovog elaborata. Grafički prilozi (prilog B) prikazuju krivinu, uzdužni profil i profile brzina po segmentima trase za sve smjerove vožnje. Za grafičke priloge u razmjeri 1:1000 podaci iz radne knjige (listovi *Profil_n*) mogu se direktno prenijeti u CAD; prilozi u ovom dokumentu prilagođeni su formatu A4.\n'))
    app = [(c, f"A{i + 1}") for i, c in enumerate(cases)]
    A("**Prilozi:** " + "; ".join(f"{n} – {c.key}: {c.label}" for c, n in app) + _T('; B – Grafički prilozi; C – Pretpostavke i ograničenja.'))
    A("\n\nZZSEGA\n\n")
    A(PB)
    for c, an in app:
        A(_TF('# Prilog {0} – Numerički proračun profila brzina, {1}: {2} {{.unnumbered}}\n', an, c.key, c.label))
        A(_TF('Vrijednosti na svakoj 10. tački mreže (korak {0} m; {1} rastuće stacionaže; vremena i gorivo računati od ulaza u dionicu); potpuni proračun na koraku {2} m nalazi se u radnoj knjizi (listovi *Profil_{3}* i *Voz1_{4}* …). Oznake: V~p~ – putničko vozilo; ', fm(P.korak * 10, 0), _T('u smjeru') if c.sg == 1 else _T('suprotno smjeru'), fm(P.korak, 0), c.key, c.key) + ", ".join(f"V~{i + 1}~ – {v['kratko'].lower()}" for i, v in enumerate(P.vozila)) + _T('; t – kumulativno vrijeme prvog teškog vozila [s]; B – kumulativna potrošnja prvog teškog vozila [l].\n\n'))
        rows = []
        for i in list(range(0, c.N, 10)) + [c.N - 1]:
            if rows and fm(c.S[i], 0) == rows[-1][0]: continue
            e = c.axis.elems[c.IDX[i]]; Rt = "∞" if c.RR[i] == 0 else fm(c.RR[i], 0)
            rows.append([fm(c.S[i], 0), fm(c.Z[i], 2), fm(c.G[i], 2), {"PRAVAC": "P", "KRUZNI_LUK": "K", "PRIJELAZNICA": "Pr"}[e["typ"]], Rt, fm(c.V[i], 1)] + [fm(t["v"][i], 1) for t in c.trk] + [fm(c.trk[0]["time"][i], 0), fm(c.trk[0]["fuel"][i], 2)])
        A(tbl(["s [m]", "z [m]", "i [%]", "El.", "R [m]", "V~p~"] + [f"V~{i + 1}~" for i in range(len(P.vozila))] + ["t [s]", "B [l]"], rows))
        A(PB)
    REG.append(("Prilozi A", _T('s – stacionaža; z – kota; i – nagib; El. – P pravac, K kružni luk, Pr prijelaznica; R; V~p~ – putničko; V~1~, V~2~ … – brzine teških vozila; t – kumulativno vrijeme [s]; B – kumulativna potrošnja [l], računato od ulaza u dionicu.')))
    A("\n\nZZSEGB\n\n")
    A(_T('# Prilog B – Grafički prilozi {.unnumbered}\n'))
    A(_TF('Svaki prilog prikazuje segment trase: krivinu 1000/R, nivelete i profile brzina za sva mjerodavna vozila (po smjerovima vožnje). Isprekidana crvena linija – računska brzina V~r~ = {0} km/h.\n', fm(Vr, 0)))
    for k, p in enumerate(F["strips"]): A(f"![]({os.path.basename(p)}){{width=15.5cm}}\n\n")
    A(PB)
    A("\n\nZZSEGC\n\n")
    A(_T('# Prilog C – Pretpostavke i ograničenja {.unnumbered}\n'))
    pre = [
        _TF('**Zakonitost V~p~ – R** zamijenjena je interpolacijom tablice R~min~(V~r~) iz Priloga 2 uz ograničenje na max V~r~ = {0} km/h (dijagram: slika 3-07 „Teorijska zavisnost projektne brzine od radijusa horizontalne krivine”, Prilog 2 pravilnika [1]). Ako se iz originalnog grafika očitaju niže brzine u krivinama, profil V~p~ će biti niži i raznolikiji (veći D~h~).', fm(maxVr, 0)),
        _T('**Zakonitost V~p~ – i~N~** je tablica sastavljena na osnovu dijagrama iz slike 3-08 Priloga 2 pravilnika [1], [4].'),
        _TF('**Ubrzanje i usporenje** a = d = {0} m/s² (ekvivalent obrasca L = 0,1·V~sr~·ΔV iz slike 3-09 Priloga 2, nominalno 0,8 m/s²).', fm(a_, 2)),
        _T('**Prijelaznice** se tretiraju kao „pravac” u smislu ciljne brzine (brzina krivine primjenjuje se na kružnom luku), što je uobičajeno pri formiranju profila V85.'),
        _T('**Dinamički model vozila**: svi parametri vozila (putničko i teška; masa, snaga, C~d~A, v~lim~, η, γ~m~, μ, specifična potrošnja, prazan hod) su pretpostavljeni; početna brzina na ulazu u trasu = v~lim~. Potrošnja je modelska i ne zamjenjuje mjerenja.'),
        _T('**Smjerovi i lanac**: dionice u smjeru rastuće stacionaže računaju se u tom smjeru, a u suprotnom smjeru obrnuto; na granicama sekcija brzina putničkog i teških vozila prenosi se s prethodne dionice u smjeru vožnje; ulazna brzina putničkog vozila na početku lanca iznosi ') + fm(P.ulazna_brzina, 0) + _T(' km/h. Pretpostavlja se da se vozilo na granici sekcija kreće bez zaustavljanja, a stacionaže osovina u susjednim sekcijama uzimaju se kao nominalno jednake.'),
        _T('**Granice dionica**: lukovi koji prelaze granicu sekcije ocjenjuju se samo u dijelu unutar sekcije.'),
        _T('**Poprečni nagib** primijenjenog luka = nagib vanjske trake kolovoza lijeve strane (desna je jednaka sa suprotnim predznakom); nagib se koristi samo u provjerama (tabele poprečnih nagiba), ne i u profilu brzine, vremenu i gorivu.'),
        _T('**Kriterij ΔV** (10 / 20 km/h) preuzet je iz literature (Lamm i dr.) [5] i primijenjen na razliku ciljnih brzina elemenata; ako Priručnik ima drugačije granice, uskladiti.'),
        _T('**Preglednost** (Elaborat 4) nije obrađena.'),
    ]
    pre.append(_TF('**Referentni najveći podužni nagib** za V~r~ = {0} km/h uzet je {1} % (potvrditi).', fm(Vr, 0), fm(P.max_nagib, 1)))
    for i, t in enumerate(pre, 1): A(f"{i}. {t}")
    if cfg.get("napomene"):
        nn = cfg["napomene"]; nn = nn if isinstance(nn, list) else [nn]
        for t in nn: A(f"{len(pre) + 1 + nn.index(t)}. {t}")
    A("\n# Literatura {.unnumbered}\n")
    A("Brojevi odgovaraju oznakama [n] u tekstu.\n\n")
    A(tbl(["Br.", "Izvor", _T('Šta je preuzeto / gdje se koristi')], [[n + 1, a, b] for n, (a, b) in enumerate(tr_obj(LIT))], ["c", "l", "l"]))
    # TABCOL
    tc = _T('**Tabela O2 – Objašnjenje kolona u tabelama**\n\n') + tbl([_T('Tabela'), _T('Značenje kolona')], [list(r) for r in REG], ["l", "l"])
    global LAST_REG
    LAST_REG = REG
    txt = "\n\n".join(md).replace("@@TABCOL@@", tc)
    txt = re.sub(r"\$\$.*?\$\$", lambda m: re.sub(r"(\d),(\d)", r"\1{,}\2", m.group(0)), txt, flags=re.S)
    return txt


def run_pandoc(md_path, out_path, resource_dir, ref):
    from .paths import find_pandoc
    exe = find_pandoc()
    cmd = [exe, md_path, "-f", "markdown+pipe_tables+tex_math_dollars+subscript+raw_attribute", "-o", out_path, f"--reference-doc={ref}", "--number-sections", "--toc", "--toc-depth=2", f"--resource-path={resource_dir}"]
    r = subprocess.run(cmd, capture_output=True, text=True, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    if r.returncode != 0: raise RuntimeError(_T('Pandoc greška: ') + r.stderr[:2000])


def postprocess_tables(tbl_elems):
    from docx.table import Table
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tables = [Table(e, None) for e in tbl_elems]
    def borders(tb):
        tblPr = tb._tbl.tblPr
        for e in tblPr.findall(qn('w:tblBorders')): tblPr.remove(e)
        b = OxmlElement('w:tblBorders')
        for n in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            x = OxmlElement('w:' + n); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:space'), '0'); x.set(qn('w:color'), '808080'); b.append(x)
        tblPr.append(b)
        for e in tblPr.findall(qn('w:tblW')): tblPr.remove(e)
        w = OxmlElement('w:tblW'); w.set(qn('w:w'), '5000'); w.set(qn('w:type'), 'pct'); tblPr.append(w)
        lay = tblPr.find(qn('w:tblLayout'))
        if lay is None: lay = OxmlElement('w:tblLayout'); tblPr.append(lay)
        lay.set(qn('w:type'), 'autofit')
        mar = OxmlElement('w:tblCellMar')
        for n in ('left', 'right'):
            x = OxmlElement('w:' + n); x.set(qn('w:w'), '60'); x.set(qn('w:type'), 'dxa'); mar.append(x)
        tblPr.append(mar)
    for t in tables:
        borders(t); ncol = len(t.columns); fs = 8 if ncol >= 9 else 9
        for ri, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
            if ri == 0: trPr.append(OxmlElement('w:tblHeader'))
            for c in row.cells:
                tcPr = c._tc.get_or_add_tcPr()
                for e in tcPr.findall(qn('w:tcW')): tcPr.remove(e)
                w = OxmlElement('w:tcW'); w.set(qn('w:w'), '0'); w.set(qn('w:type'), 'auto'); tcPr.insert(0, w)
                if ri == 0:
                    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'D9E2F3'); tcPr.append(sh)
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(fs); r.font.name = "Arial"
                        if ri == 0: r.font.bold = True
        for gc in t._tbl.tblGrid.findall(qn('w:gridCol')): gc.set(qn('w:w'), str(int(9000 / max(ncol, 1))))


def postprocess(raw, out, footer):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = Document(raw)
    postprocess_tables([t._tbl for t in d.tables])
    sec = d.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7); sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2.2)
    fp = sec.footer.paragraphs[0]; fp.alignment = 1
    def fld(p, code):
        r = p.add_run(); a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin'); r._r.append(a)
        r = p.add_run(); i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = code; r._r.append(i)
        r = p.add_run(); a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'end'); r._r.append(a)
    fp.add_run(footer + " | strana ")
    fld(fp, "PAGE")
    for r in fp.runs: r.font.size = Pt(8); r.font.name = "Arial"
    for p in d.paragraphs:
        if p.style.name in ("Captioned Figure", "Image Caption"): p.alignment = 1
    ps = d.paragraphs
    for pp in ps:
        if pp.text.startswith(_T('Tabela ')) and len(pp.text) < 200: pp.paragraph_format.keep_with_next = True
        if pp.style.name == "TOC Heading" or pp.text.strip() == _T('Sadržaj'): pp.paragraph_format.page_break_before = True
    h1 = [pp for pp in ps if pp.style.name == "Heading 1"]
    if h1: h1[0].paragraph_format.page_break_before = True
    st = d.settings.element
    u = OxmlElement('w:updateFields'); u.set(qn('w:val'), 'true'); anchor = None
    for tag in ('w:hdrShapeDefaults', 'w:footnotePr', 'w:endnotePr', 'w:compat', 'w:docVars', 'w:rsids', 'm:mathPr', 'w:themeFontLang', 'w:clrSchemeMapping', 'w:decimalSymbol', 'w:listSeparator'):
        try: f = st.find(qn(tag))
        except Exception: f = None
        if f is not None: anchor = f; break
    if anchor is not None: anchor.addprevious(u)
    else: st.append(u)
    z = st.find(qn('w:zoom'))
    if z is not None: st.remove(z); st.insert(0, z)
    from .template_export import normalize_order
    normalize_order(d.element)
    d.save(out)


def write_docx(P, sections, cases, axes, fw, bw, has_caps, F, cfg, out_docx, log=print):
    txt = build_report(P, sections, cases, axes, fw, bw, has_caps, F, cfg, os.path.dirname(out_docx), log)
    txt = re.sub(r"\n*ZZSEG[ABC]\n*", "\n\n", txt)
    with tempfile.TemporaryDirectory() as td:
        figdir = os.path.join(td, "fig")
        os.makedirs(figdir)
        for p in [F[k] for k in F if k != "strips"] + F["strips"]: shutil.copy(p, figdir)
        txt = txt.replace("](", "](fig/") if False else re.sub(r"\]\(([^)/]+\.png)\)", r"](fig/\1)", txt)
        mdp = os.path.join(td, "doc.md"); open(mdp, "w", encoding="utf-8").write(txt)
        open(os.path.join(os.path.dirname(out_docx), "elaborat3_izvor.md"), "w", encoding="utf-8").write(txt) if cfg.get("zadrzi_md") else None
        raw = os.path.join(td, "raw.docx")
        run_pandoc(mdp, raw, td, os.path.join(HERE, "data", "ref.docx"))
        postprocess(raw, out_docx, _TF('Elaborat 3 – Vozno-dinamičke analize | {0}', cfg.get('projekat', '')))
    log(_TF("Word: {0}", out_docx))


def write_docx_template(P, sections, cases, axes, fw, bw, has_caps, F, cfg, out_docx, log=print):
    from . import template_export
    txt = build_report(P, sections, cases, axes, fw, bw, has_caps, F, cfg, os.path.dirname(out_docx), log, template_mode=True)
    with tempfile.TemporaryDirectory() as td:
        figdir = os.path.join(td, "fig"); os.makedirs(figdir)
        for p in [F[k] for k in F if k != "strips"] + F["strips"]: shutil.copy(p, figdir)
        txt = re.sub(r"\]\(([^)/]+\.png)\)", r"](fig/\1)", txt)
        mdp = os.path.join(td, "doc.md"); open(mdp, "w", encoding="utf-8").write(txt)
        raw = os.path.join(td, "raw.docx")
        from .paths import find_pandoc
        exe = find_pandoc()
        cmd = [exe, mdp, "-f", "markdown+pipe_tables+tex_math_dollars+subscript+raw_attribute", "-o", raw, f"--reference-doc={os.path.join(HERE, 'data', 'ref.docx')}", "--number-sections", f"--resource-path={td}"]
        rr = subprocess.run(cmd, capture_output=True, text=True, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        if rr.returncode != 0: raise RuntimeError(_T('Pandoc greška: ') + rr.stderr[:2000])
        template_export.export(raw, cfg["template"], out_docx, log)
